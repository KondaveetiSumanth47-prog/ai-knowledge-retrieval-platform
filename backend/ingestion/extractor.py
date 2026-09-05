import os
import pymupdf as fitz  # PyMuPDF
import docx
import pandas as pd

class DocumentExtractor:
    """
    Extracts text from PDF, DOCX, TXT, and CSV file formats.
    """
    
    @staticmethod
    def extract_from_pdf(file_path: str) -> str:
        """Extract text from PDF using PyMuPDF (fitz)"""
        text_content = []
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text.strip():
                text_content.append(f"--- Page {page_num + 1} ---\n" + page_text.strip())
        doc.close()
        return "\n\n".join(text_content)

    @staticmethod
    def extract_from_docx(file_path: str) -> str:
        """Extract text and tables from DOCX using python-docx"""
        doc = docx.Document(file_path)
        text_content = []
        
        # Extract paragraph text
        for p in doc.paragraphs:
            if p.text.strip():
                text_content.append(p.text.strip())
                
        # Extract table text
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                text_content.append(f"--- Table {table_idx + 1} ---\n" + "\n".join(table_rows))
                
        return "\n\n".join(text_content)

    @staticmethod
    def extract_from_txt(file_path: str) -> str:
        """Extract text from plain text file with multi-encoding fallback"""
        encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode text file {file_path} with supported encodings.")

    @staticmethod
    def extract_from_csv(file_path: str) -> str:
        """Extract text from CSV using pandas, formatting rows as structured key-value passages"""
        df = pd.read_csv(file_path)
        rows_text = []
        
        # Overview header
        rows_text.append(f"CSV Dataset Overview: {len(df)} records, columns: {', '.join(df.columns)}")
        
        for idx, row in df.iterrows():
            record_str = f"Record #{idx + 1}: " + ", ".join([f"{col}: {val}" for col, val in row.items() if pd.notna(val)])
            rows_text.append(record_str)
            
        return "\n".join(rows_text)

    @classmethod
    def extract(cls, file_path: str, file_type: str = None) -> str:
        """
        Unified extraction method matching file extension or specified file_type.
        """
        if not file_type:
            ext = os.path.splitext(file_path)[1].lower().strip('.')
            file_type = ext

        file_type = file_type.lower()
        if file_type == 'pdf':
            return cls.extract_from_pdf(file_path)
        elif file_type in ['docx', 'doc']:
            return cls.extract_from_docx(file_path)
        elif file_type in ['txt', 'text', 'md']:
            return cls.extract_from_txt(file_path)
        elif file_type == 'csv':
            return cls.extract_from_csv(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_type}")
